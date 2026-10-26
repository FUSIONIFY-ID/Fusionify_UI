"""FUSIONIFY UI · template Excel (.xlsx).

  python3 scripts/dokumen/sheet.py   → dokumen/office/FUSIONIFY-Spreadsheet.xlsx

Sheet: Panduan, Invoice (berformula), Anggaran, Daftar (pelacak dengan status kit).
Sel isian berlatar kuning muda (--beta-bg), sama seperti teks "ganti" di template lain.
Warna dan huruf mengikuti css/tokens.css. Huruf: pasang tokens/fonts-ttf/*.ttf; kalau belum,
Excel memakai huruf pengganti dan isinya tetap terbaca.
"""
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
ASET = ROOT / "dokumen" / "office" / "aset"
OUT = ROOT / "dokumen" / "office" / "FUSIONIFY-Spreadsheet.xlsx"

# Token (css/tokens.css)
BLUE, DEEP, LIGHT, INK, INK2, INK3 = "0D6CD7", "0261CD", "539DE9", "0B0F1A", "3A4458", "566175"
SURFACE, PAPER, LINE, WHITE = "EAF4FF", "F7FBFF", "D3E4F7", "FFFFFF"
STATUS = [  # (teks, warna teks, latar) = chip go, beta, wip, plan, eos
    ("Aktif", "157F3B", "E3F6E9"),
    ("Beta", "9A5B00", "FFF2D8"),
    ("Dalam pengembangan", "8A4A00", "FFE9D6"),
    ("Direncanakan", "0261CD", "EEF7FF"),
    ("Layanan berakhir", "5B6475", "ECEFF4"),
]
ISIAN = "FFF2D8"     # --beta-bg: sel yang diisi pengguna
DANGER, DANGER_BG = "B4232C", "FDE8EA"

BODY, HEAD, NUM = "Sansation", "Afacad Flux ExtraBold", "Exo 2"
RUPIAH = '"Rp"#,##0;-"Rp"#,##0;"-"'
PERSEN = '0.0%;-0.0%;"-"'
TANGGAL = '[$-421]d mmmm yyyy'  # nama bulan bahasa Indonesia

f_body = Font(name=BODY, size=10, color=INK)
f_muted = Font(name=BODY, size=9, color=INK3)
f_bold = Font(name=BODY, size=10, bold=True, color=INK)
f_kicker = Font(name=NUM, size=8, bold=True, color=DEEP)
f_head_tbl = Font(name=NUM, size=9, bold=True, color=WHITE)
f_judul = Font(name=HEAD, size=24, color=INK)
f_total = Font(name=HEAD, size=13, color=WHITE)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
tipis = Side(style="thin", color=LINE)
tebal = Side(style="medium", color=INK)
garis_bawah = Border(bottom=tipis)
kiri = Alignment(horizontal="left", vertical="center", wrap_text=True)
kanan = Alignment(horizontal="right", vertical="center")
tengah = Alignment(horizontal="center", vertical="center")

wb = Workbook()


def dasar(ws, lebar):
    ws.sheet_view.showGridLines = False
    for kol, w in lebar.items():
        ws.column_dimensions[kol].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = "PT FUSIONIFY DIGITAL STORE · fusionify.id"
    ws.oddFooter.left.font = NUM
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Halaman &P dari &N"
    ws.oddFooter.right.font = NUM
    ws.oddFooter.right.size = 8


def kepala_tabel(ws, baris, judul, kolom_awal=1):
    for i, t in enumerate(judul):
        c = ws.cell(row=baris, column=kolom_awal + i, value=t.upper())
        c.font, c.fill, c.alignment = f_head_tbl, fill(INK), Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[baris].height = 22


def isian(c, nilai=None, fmt=None):
    if nilai is not None:
        c.value = nilai
    c.fill = fill(ISIAN)
    c.font = f_body
    if fmt:
        c.number_format = fmt
    return c


def logo(ws, sel="A1"):
    img = XLImage(str(ASET / "fusionify-logo-horizontal.png"))
    img.height, img.width = 44, int(44 * 1087 / 340)
    ws.add_image(img, sel)


def kop(ws, kol_info):
    logo(ws)
    for i, (t, f) in enumerate([("PT FUSIONIFY DIGITAL STORE", Font(name=HEAD, size=12, color=INK)),
                                ("Bogor, Jawa Barat", f_muted),
                                ("hello@fusionify.id · +62 831-1138-0628 · fusionify.id", f_muted)]):
        c = ws[f"{kol_info}{i + 1}"]
        c.value, c.font, c.alignment = t, f, Alignment(horizontal="right")
    for r in range(1, 4):
        ws.row_dimensions[r].height = 16


# ---------------- Panduan ----------------
ws = wb.active
ws.title = "Panduan"
dasar(ws, {"A": 3, "B": 26, "C": 70})
logo(ws, "B2")
ws["B6"].value, ws["B6"].font = "TEMPLATE SPREADSHEET FUSIONIFY", f_kicker
ws["B7"].value, ws["B7"].font = "Cara memakai file ini", Font(name=HEAD, size=20, color=INK)
ws.row_dimensions[7].height = 28
baris = [
    ("Sel kuning muda", "Isian. Ganti nilainya dengan data asli. Hanya sel ini yang perlu diubah.", ISIAN),
    ("Sel putih", "Rumus. Jangan diketik ulang; nilainya dihitung otomatis.", WHITE),
    ("Baris contoh", "Setiap sheet punya satu baris contoh (ditandai di kolom catatan). Hapus atau timpa sebelum dipakai.", WHITE),
    ("Invoice", "Tagihan berformula: jumlah, subtotal, diskon, pajak, total. Isi pajak hanya kalau perusahaan sudah PKP.", WHITE),
    ("Anggaran", "Rencana vs realisasi per kategori. Persentase di atas 100% otomatis ditandai merah.", WHITE),
    ("Daftar", "Pelacak umum (produk, tugas, aset). Kolom Status memakai lima status resmi kit dengan warnanya.", WHITE),
    ("Huruf", "Pasang huruf di tokens/fonts-ttf (Afacad Flux, Sansation, Exo 2) supaya tampilan sama dengan dokumen lain.", WHITE),
    ("Jujur", "Jangan isi angka karangan. Data contoh harus diganti data asli sebelum dibagikan.", WHITE),
]
for i, (a, b, warna) in enumerate(baris):
    r = 9 + i
    ca, cb = ws[f"B{r}"], ws[f"C{r}"]
    ca.value, ca.font, ca.fill, ca.alignment, ca.border = a, f_bold, fill(warna), kiri, garis_bawah
    cb.value, cb.font, cb.alignment, cb.border = b, f_body, kiri, garis_bawah
    ws.row_dimensions[r].height = 30

# ---------------- Invoice ----------------
ws = wb.create_sheet("Invoice")
dasar(ws, {"A": 6, "B": 44, "C": 9, "D": 17, "E": 19})
kop(ws, "E")
for col in "ABCDE":
    ws[f"{col}4"].border = Border(bottom=tebal)
ws["A6"].value, ws["A6"].font = "TAGIHAN", f_kicker
ws["A7"].value, ws["A7"].font = "Invoice", f_judul
ws.row_dimensions[7].height = 34
for i, (label, nilai, fmt) in enumerate([("Nomor", "012/INV/FDS/X/2026", None), ("Tanggal", date.today(), TANGGAL), ("Jatuh tempo", "=D7+14", TANGGAL)]):
    r = 6 + i
    ws[f"C{r}"].value, ws[f"C{r}"].font, ws[f"C{r}"].alignment = label, f_muted, kanan
    c = ws[f"D{r}"]
    ws.merge_cells(f"D{r}:E{r}")
    if label == "Jatuh tempo":
        c.value, c.font, c.number_format, c.alignment = "=D7+14", f_body, TANGGAL, kanan
        c.comment = Comment("Rumus: tanggal + 14 hari. Ganti angka 14 kalau termin berbeda.", "FUSIONIFY")
    else:
        isian(c, nilai, fmt).alignment = kanan
        if label == "Tanggal":
            c.comment = Comment("Tanggal tetap invoice. Jangan pakai =TODAY(): nilainya berubah setiap file dibuka.", "FUSIONIFY")

ws["A10"].value, ws["A10"].font = "DITAGIHKAN KEPADA", f_kicker
for i, t in enumerate(["Nama klien / perusahaan", "Alamat klien", "Email atau nomor telepon"]):
    ws.merge_cells(f"A{11 + i}:B{11 + i}")
    isian(ws[f"A{11 + i}"], t).font = f_bold if i == 0 else f_body
ws["D10"].value, ws["D10"].font = "PROYEK", f_kicker
ws.merge_cells("D11:E11")
isian(ws["D11"], "Nama proyek")

HDR = 15
kepala_tabel(ws, HDR, ["No", "Deskripsi", "Qty", "Harga", "Jumlah"])
for c in ("C", "D", "E"):
    ws[f"{c}{HDR}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
AWAL, AKHIR = HDR + 1, HDR + 10
for r in range(AWAL, AKHIR + 1):
    ws[f"A{r}"] = f'=IF(B{r}="","",ROW()-{HDR})'
    ws[f"A{r}"].font, ws[f"A{r}"].alignment = Font(name=NUM, size=10, color=INK3), tengah
    for col in "BCD":
        isian(ws[f"{col}{r}"])
    ws[f"B{r}"].alignment = kiri
    ws[f"C{r}"].number_format, ws[f"C{r}"].alignment = "#,##0", kanan
    ws[f"D{r}"].number_format = RUPIAH
    ws[f"E{r}"] = f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})'
    ws[f"E{r}"].number_format, ws[f"E{r}"].font = RUPIAH, Font(name=NUM, size=10, color=INK)
    for col in "ABCDE":
        ws[f"{col}{r}"].border = garis_bawah
    ws.row_dimensions[r].height = 20
# Satu baris contoh
ws[f"B{AWAL}"], ws[f"C{AWAL}"], ws[f"D{AWAL}"] = "Contoh: pembuatan website company profile", 1, 2500000
ws[f"B{AWAL}"].comment = Comment("Baris contoh. Timpa dengan pekerjaan asli.", "FUSIONIFY")

r = AKHIR + 2
ringkas = [("Subtotal", f"=SUM(E{AWAL}:E{AKHIR})"), ("Diskon", 0), ("Pajak (%)", 0), ("Pajak", None), ("Total", None)]
pos = {}
for i, (label, nilai) in enumerate(ringkas):
    rr = r + i
    pos[label] = rr
    ws[f"D{rr}"].value, ws[f"D{rr}"].font, ws[f"D{rr}"].alignment = label, f_body, kanan
pos_sub, pos_disk, pos_pct, pos_paj, pos_tot = (pos[k] for k in ["Subtotal", "Diskon", "Pajak (%)", "Pajak", "Total"])
ws[f"E{pos_sub}"] = f"=SUM(E{AWAL}:E{AKHIR})"
ws[f"E{pos_sub}"].number_format = RUPIAH
isian(ws[f"E{pos_disk}"], 0, RUPIAH)
isian(ws[f"E{pos_pct}"], 0, PERSEN)
ws[f"E{pos_pct}"].comment = Comment("Isi 0,11 untuk PPN 11% hanya kalau perusahaan sudah PKP. Selain itu biarkan 0.", "FUSIONIFY")
ws[f"E{pos_paj}"] = f"=ROUND(MAX(0,E{pos_sub}-E{pos_disk})*E{pos_pct},0)"
ws[f"E{pos_paj}"].number_format = RUPIAH
ws[f"E{pos_tot}"] = f"=MAX(0,E{pos_sub}-E{pos_disk})+E{pos_paj}"
for col in "DE":
    c = ws[f"{col}{pos_tot}"]
    c.fill, c.font = fill(BLUE), f_total
ws[f"E{pos_tot}"].number_format = RUPIAH
ws.row_dimensions[pos_tot].height = 26
for rr in (pos_sub, pos_disk, pos_pct, pos_paj):
    ws[f"E{rr}"].font = Font(name=NUM, size=10, color=INK)
    ws[f"E{rr}"].alignment = kanan

rb = pos_tot + 2
ws[f"A{rb}"].value, ws[f"A{rb}"].font = "PEMBAYARAN", f_kicker
for i, t in enumerate(["Nama bank", "No. rekening", "a.n. PT FUSIONIFY DIGITAL STORE"]):
    ws.merge_cells(f"A{rb + 1 + i}:B{rb + 1 + i}")
    c = ws[f"A{rb + 1 + i}"]
    if i < 2:
        isian(c, t)
    else:
        c.value, c.font = t, f_bold
ws.merge_cells(f"A{rb + 4}:B{rb + 4}")
ws[f"A{rb + 4}"].value, ws[f"A{rb + 4}"].font = "Cantumkan nomor invoice di berita transfer, lalu kirim bukti ke hello@fusionify.id.", f_muted
ws[f"D{rb + 1}"].value, ws[f"D{rb + 1}"].font, ws[f"D{rb + 1}"].alignment = "Hormat kami,", f_body, tengah
ws.merge_cells(f"D{rb + 1}:E{rb + 1}")
ws.merge_cells(f"D{rb + 5}:E{rb + 5}")
isian(ws[f"D{rb + 5}"], "Nama penanda tangan").alignment = tengah
ws.merge_cells(f"D{rb + 6}:E{rb + 6}")
isian(ws[f"D{rb + 6}"], "Jabatan").alignment = tengah
ws.print_area = f"A1:E{rb + 6}"

# ---------------- Anggaran ----------------
ws = wb.create_sheet("Anggaran")
dasar(ws, {"A": 6, "B": 32, "C": 18, "D": 18, "E": 18, "F": 13, "G": 34})
kop(ws, "G")
for col in "ABCDEFG":
    ws[f"{col}4"].border = Border(bottom=tebal)
ws["A6"].value, ws["A6"].font = "KEUANGAN", f_kicker
ws["A7"].value, ws["A7"].font = "Anggaran", f_judul
ws.row_dimensions[7].height = 34
ws["E6"].value, ws["E6"].font, ws["E6"].alignment = "Periode", f_muted, kanan
ws.merge_cells("F6:G6")
isian(ws["F6"], "Oktober 2026")
HDR = 9
kepala_tabel(ws, HDR, ["No", "Kategori", "Rencana", "Realisasi", "Selisih", "Terpakai", "Catatan"])
AWAL, AKHIR = HDR + 1, HDR + 15
for r in range(AWAL, AKHIR + 1):
    ws[f"A{r}"] = f'=IF(B{r}="","",ROW()-{HDR})'
    ws[f"A{r}"].font, ws[f"A{r}"].alignment = Font(name=NUM, size=10, color=INK3), tengah
    for col, fmt in (("B", None), ("C", RUPIAH), ("D", RUPIAH), ("G", None)):
        isian(ws[f"{col}{r}"], fmt=fmt)
    ws[f"E{r}"] = f'=IF(B{r}="","",C{r}-D{r})'
    ws[f"F{r}"] = f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}/C{r})'
    ws[f"E{r}"].number_format, ws[f"F{r}"].number_format = RUPIAH, PERSEN
    for col in "EF":
        ws[f"{col}{r}"].font = Font(name=NUM, size=10, color=INK)
    for col in "ABCDEFG":
        ws[f"{col}{r}"].border = garis_bawah
ws[f"B{AWAL}"], ws[f"C{AWAL}"], ws[f"D{AWAL}"], ws[f"G{AWAL}"] = "Contoh: server dan domain", 1500000, 1250000, "Baris contoh, timpa dengan data asli"
T = AKHIR + 1
ws[f"B{T}"] = "Total"
ws[f"C{T}"], ws[f"D{T}"] = f"=SUM(C{AWAL}:C{AKHIR})", f"=SUM(D{AWAL}:D{AKHIR})"
ws[f"E{T}"] = f"=C{T}-D{T}"
ws[f"F{T}"] = f'=IF(C{T}=0,"",D{T}/C{T})'
for col in "ABCDEFG":
    c = ws[f"{col}{T}"]
    c.fill, c.font = fill(BLUE), Font(name=HEAD if col == "B" else NUM, size=11, color=WHITE, bold=col != "B")
for col in "CDE":
    ws[f"{col}{T}"].number_format = RUPIAH
ws[f"F{T}"].number_format = PERSEN
ws.row_dimensions[T].height = 24
# Melebihi rencana → merah (--danger)
ws.conditional_formatting.add(f"F{AWAL}:F{AKHIR}", CellIsRule(operator="greaterThan", formula=["1"], font=Font(name=NUM, color=DANGER, bold=True), fill=fill(DANGER_BG)))
ws.conditional_formatting.add(f"E{AWAL}:E{AKHIR}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=NUM, color=DANGER, bold=True)))
ws.freeze_panes = f"A{AWAL}"
ws.print_area = f"A1:G{T}"

# ---------------- Daftar ----------------
ws = wb.create_sheet("Daftar")
dasar(ws, {"A": 6, "B": 30, "C": 18, "D": 22, "E": 22, "F": 16, "G": 38})
kop(ws, "G")
for col in "ABCDEFG":
    ws[f"{col}4"].border = Border(bottom=tebal)
ws["A6"].value, ws["A6"].font = "PELACAK", f_kicker
ws["A7"].value, ws["A7"].font = "Daftar", f_judul
ws.row_dimensions[7].height = 34
HDR = 9
kepala_tabel(ws, HDR, ["No", "Nama", "Kategori", "Status", "Penanggung jawab", "Tanggal", "Catatan"])
AWAL, AKHIR = HDR + 1, HDR + 50
dv = DataValidation(type="list", formula1='"' + ",".join(s[0] for s in STATUS) + '"', allow_blank=True,
                    showErrorMessage=True, errorTitle="Status tidak dikenal", error="Pilih salah satu status resmi FUSIONIFY.")
ws.add_data_validation(dv)
dv.add(f"D{AWAL}:D{AKHIR}")
for r in range(AWAL, AKHIR + 1):
    ws[f"A{r}"] = f'=IF(B{r}="","",ROW()-{HDR})'
    ws[f"A{r}"].font, ws[f"A{r}"].alignment = Font(name=NUM, size=10, color=INK3), tengah
    for col in "BCDEFG":
        isian(ws[f"{col}{r}"])
    ws[f"F{r}"].number_format = TANGGAL
    for col in "ABCDEFG":
        ws[f"{col}{r}"].border = garis_bawah
ws[f"B{AWAL}"], ws[f"C{AWAL}"], ws[f"D{AWAL}"], ws[f"E{AWAL}"], ws[f"F{AWAL}"], ws[f"G{AWAL}"] = (
    "Contoh: website company profile", "Proyek klien", "Dalam pengembangan", "Nama orang", date.today(), "Baris contoh, timpa dengan data asli")
for teks, warna, latar in STATUS:
    ws.conditional_formatting.add(f"D{AWAL}:D{AKHIR}", FormulaRule(formula=[f'$D{AWAL}="{teks}"'], font=Font(name=NUM, bold=True, color=warna), fill=fill(latar)))
ws.auto_filter.ref = f"A{HDR}:G{AKHIR}"
ws.freeze_panes = f"A{AWAL}"
ws.print_title_rows = f"{HDR}:{HDR}"

wb.properties.creator = "PT FUSIONIFY DIGITAL STORE"
wb.calculation.fullCalcOnLoad = True  # Excel menghitung ulang semua rumus saat file dibuka
wb.properties.title = "Template spreadsheet FUSIONIFY"
wb.active = wb.sheetnames.index("Invoice")
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"✓ {OUT.relative_to(ROOT)}")
